import sys
import os
import shutil
import datetime


sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))
from MACD.MACD import MACDPlot
from parse.TimesAlgorithm import TimesAlgorithm
from db.db import MSSQLDbContext



if __name__ == '__main__':
    import pandas as pd
    import random
    excel_file = 'MACD.xlsx'  # 新增，參考 TimesMACDFeed.py
    _ExcelPath = r"C:\Programs\test_data"
    sql = MSSQLDbContext({'server': 'wpdb2.hihosting.hinet.net', 'user': 'p89880749_p89880749',
                          'password': 'Jonny1070607!@#$%', 'database': 'p89880749_test'})
    rows = sql.select('select drawTerm, bigShowOrder from Bingo ORDER BY drawTerm ASC ')
    # 增加 bigShowOrders 欄位
    for row in rows:
        if 'bigShowOrder' in row and isinstance(row['bigShowOrder'], str):
            row['bigShowOrders'] = row['bigShowOrder'].split(',')
        else:
            row['bigShowOrders'] = []
    algo = TimesAlgorithm()
    fileNameNoExten = os.path.splitext(os.path.basename(excel_file))[0]
    fileExten = os.path.splitext(os.path.basename(excel_file))[1]
    print(fileNameNoExten)
    print(fileExten)
    # takeColumns 設定為 1~80 共 80 個元素
    takeColumns = [str(i) for i in range(1, 81)]
    algo.TakeColumns = takeColumns
    # 以 rows['bigShowOrders'] 轉為 list 作為 inputs
    inputs = list(row['bigShowOrders'] for row in rows if 'bigShowOrders' in row)
    algo.LoadData(inputs)
    # 建立一個 df 等於 algo.DfResult，最開頭插入一欄 drawTerm
    df = algo.DfResult.copy()
    if isinstance(rows, list) and len(rows) > 0 and 'drawTerm' in rows[0]:
        draw_terms = [row['drawTerm'] for row in rows]
        df.insert(0, 'drawTerm', draw_terms)
    # print(df)
    # 匯出 df 到 Excel
    export_path = os.path.join(_ExcelPath, f"export.xlsx")  # 移除
    df.to_excel(export_path, index=False)  # 移除
    plot_rows = df.to_dict(orient='records')
    balls = []
    macd_ranges = {
        'MACD <0.05': [],
        'MACD 0.05-0.1': [],
        'MACD 0.1-0.15': [],
        'MACD 0.15-0.2': [],
        'MACD >0.2': []
    }
    for j in range(1, 81):
        ball_field = str(j).zfill(2)
        import matplotlib.pyplot as plt
        macd = MACDPlot(plot_rows, date_field='drawTerm', close_field=ball_field)
        macd.load()
        offset = abs(macd.MACDOffset)
        if macd.isMACDBig:
            if offset < 0.05:
                macd_ranges['MACD <0.05'].append(ball_field)
            elif offset < 0.1:
                macd_ranges['MACD 0.05-0.1'].append(ball_field)
            elif offset < 0.15:
                macd_ranges['MACD 0.1-0.15'].append(ball_field)
            elif offset < 0.2:
                macd_ranges['MACD 0.15-0.2'].append(ball_field)
            else:
                macd_ranges['MACD >0.2'].append(ball_field)
            balls.append(ball_field)
    print(f"MACD big balls: {balls}")

    # 輸出 MACD 統計到 Sheet2，球號以逗號分隔
    macd_stats_df = pd.DataFrame({k: [','.join(v)] for k, v in macd_ranges.items()})

    # 讀取 actual.xlsx，並寫入新一列
    import pandas as pd
    from openpyxl.utils import get_column_letter
    actual_path = os.path.join(_ExcelPath, "actual.xlsx")
    from openpyxl import load_workbook
    import os as _os

    # 取得當前 drawTerm 最大值
    if not df.empty and 'drawTerm' in df.columns:
        last_draw_term = df['drawTerm'].iloc[-1]
        next_draw_term = int(last_draw_term) + 1
    else:
        next_draw_term = 1
    # 新增一列
    # 新增 bigSmallSort 欄位
    small_count = sum(1 for b in balls if 1 <= int(b) <= 40)
    big_count = sum(1 for b in balls if 41 <= int(b) <= 80)
    total = len(balls)
    smallPercent = small_count / total if total > 0 else 0
    bigPercent = big_count / total if total > 0 else 0
    if total > 0:
        if bigPercent >= 0.65:
            bigSmallSort = "大"
        elif smallPercent >= 0.65:
            bigSmallSort = "小"
        else:
            bigSmallSort = "合"
    else:
        bigSmallSort = ""
    new_row = {"drawTerm": next_draw_term, "bigShowOrder": [], "balls": balls, "bigSmallSort": bigSmallSort, "smallPercent": smallPercent, "bigPercent": bigPercent}
    new_df = pd.DataFrame([new_row])

    # 新增：DeferAlgorithm 實體化並處理
    from parse.DeferAlgorithm import DeferAlgorithm
    defer_algo = DeferAlgorithm()
    defer_algo.TakeColumns = takeColumns
    defer_algo.LoadData(inputs)
    df_defer = defer_algo.DfResult
    # 只取最後一列
    df_defer_last = df_defer.tail(1)

    # 新增：ContinAlgorithm 實體化並處理
    from parse.ContinAlgorithm import ContinAlgorithm
    contin_algo = ContinAlgorithm()
    contin_algo.TakeColumns = takeColumns
    contin_algo.LoadData(inputs)
    df_contin = contin_algo.DfResult
    # 只取最後一列
    df_contin_last = df_contin.tail(1)

    # 新增：TimesAlgorithm 實體化並處理
    times_algo = TimesAlgorithm()
    times_algo.TakeColumns = takeColumns
    times_algo.LoadData(inputs)
    df_times = times_algo.DfResult
    # 只取最後一列
    df_times_last = df_times.tail(1)

    if _os.path.exists(actual_path):
        # 檔案存在，附加寫入
        with pd.ExcelWriter(actual_path, engine='openpyxl', mode='a', if_sheet_exists='overlay') as writer:
            book = writer.book
            if "Sheet1" in book.sheetnames:
                sheet = book["Sheet1"]
                startrow = sheet.max_row
            else:
                startrow = 0
            # 只有在第一列時才寫 header
            new_df.to_excel(writer, sheet_name="Sheet1", index=False, header=(startrow==0), startrow=startrow)
            # Sheet2 也用附加方式，且標題固定
            if "Sheet2" in book.sheetnames:
                sheet2 = book["Sheet2"]
                startrow2 = sheet2.max_row
                macd_stats_df.to_excel(writer, sheet_name='Sheet2', index=False, header=(startrow2==0), startrow=startrow2)
            else:
                startrow2 = 0
                macd_stats_df.to_excel(writer, sheet_name='Sheet2', index=False, header=True, startrow=startrow2)
            # Sheet3 也用附加方式，且標題固定，改名為 Defer
            if "Defer" in book.sheetnames:
                sheet3 = book["Defer"]
                startrow3 = sheet3.max_row
                df_defer_last.to_excel(writer, sheet_name='Defer', index=False, header=(startrow3==0), startrow=startrow3)
            else:
                startrow3 = 0
                df_defer_last.to_excel(writer, sheet_name='Defer', index=False, header=True, startrow=startrow3)
            # Sheet4 也用附加方式，且標題固定，改名為 Contin
            if "Contin" in book.sheetnames:
                sheet4 = book["Contin"]
                startrow4 = sheet4.max_row
                df_contin_last.to_excel(writer, sheet_name='Contin', index=False, header=(startrow4==0), startrow=startrow4)
            else:
                startrow4 = 0
                df_contin_last.to_excel(writer, sheet_name='Contin', index=False, header=True, startrow=startrow4)
            # Sheet5 也用附加方式，且標題固定，改名為 Times
            if "Times" in book.sheetnames:
                sheet5 = book["Times"]
                startrow5 = sheet5.max_row
                df_times_last.to_excel(writer, sheet_name='Times', index=False, header=(startrow5==0), startrow=startrow5)
            else:
                startrow5 = 0
                df_times_last.to_excel(writer, sheet_name='Times', index=False, header=True, startrow=startrow5)
        # 自適應欄寬
        wb = load_workbook(actual_path)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for col in ws.columns:
                max_length = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    try:
                        if cell.value:
                            max_length = max(max_length, len(str(cell.value)))
                    except:
                        pass
                ws.column_dimensions[col_letter].width = max_length + 2
        wb.save(actual_path)
        # 直接複製 actual.xlsx 為 actual__HHmmss.xlsx
        now_str = datetime.datetime.now().strftime("%H%M%S")
        actual_view_path = os.path.join(_ExcelPath, f"actual__{now_str}.xlsx")
        shutil.copyfile(actual_path, actual_view_path)
    else:
        # 檔案不存在，建立新檔案
        with pd.ExcelWriter(actual_path, engine='openpyxl', mode='w') as writer:
            new_df.to_excel(writer, sheet_name="Sheet1", index=False, header=True)
            # 新增：輸出 MACD 統計到 Sheet2
            macd_stats_df.to_excel(writer, sheet_name='Sheet2', index=False, header=True)
            # Sheet3 輸出 DeferAlgorithm 結果（只取最後一列）
            df_defer_last.to_excel(writer, sheet_name='Defer', index=False, header=True)
            # Sheet4 輸出 ContinAlgorithm 結果（只取最後一列）
            df_contin_last.to_excel(writer, sheet_name='Contin', index=False, header=True)
            # Sheet5 輸出 TimesAlgorithm 結果（只取最後一列）
            df_times_last.to_excel(writer, sheet_name='Times', index=False, header=True)
        # 自適應欄寬
        wb = load_workbook(actual_path)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for col in ws.columns:
                max_length = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    try:
                        if cell.value:
                            max_length = max(max_length, len(str(cell.value)))
                    except:
                        pass
                ws.column_dimensions[col_letter].width = max_length + 2
        wb.save(actual_path)
        # 直接複製 actual.xlsx 為 actual__HHmmss.xlsx
        now_str = datetime.datetime.now().strftime("%H%M%S")
        actual_view_path = os.path.join(_ExcelPath, f"actual__{now_str}.xlsx")
        shutil.copyfile(actual_path, actual_view_path)

